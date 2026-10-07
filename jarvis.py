"""
╔══════════════════════════════════════════════════════════════════════════╗
║      JARVIS AI v4.0 - Tamil + English Mixed (Tanglish) Assistant         ║
║      Full Windows Control | File Creator | WhatsApp Messenger            ║
╚══════════════════════════════════════════════════════════════════════════╝

SPEAK IN TANGLISH (Tamil + English mix):
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
 APPS         : "Chrome open பண்ணு" / "open notepad"
 VOLUME       : "volume up பண்ணு" / "ஒலி அதிகரி"
 BRIGHTNESS   : "brightness up" / "பிரகாசம் அதிகரி"
 SCREENSHOT   : "screenshot எடு" / "screen capture பண்ணு"
 FILE CREATE  : "text file create பண்ணு MyFile"
               "word file create பண்ணு Report"
               "python file create பண்ணு script"
 WHATSAPP     : "whatsapp-ல அம்மாக்கு message அனுப்பு: வீட்டில் இருக்கேன்"
               "send whatsapp to friend: Hello da"
 BATTERY      : "battery status சொல்லு"
 SYSTEM INFO  : "system info சொல்லு" / "RAM எவ்வளவு"
 WINDOW CTRL  : "window close பண்ணு" / "minimize பண்ணு"
 LOCK/SLEEP   : "screen lock பண்ணு" / "sleep mode போ"
 SCROLL       : "scroll down பண்ணு" / "scroll up"
 MEDIA        : "music play பண்ணு" / "next song போ"
 YOUTUBE      : "youtube open பண்ணு"
 SEARCH       : "google-ல AI search பண்ணு"
 TIME/DATE    : "time என்ன" / "today date சொல்லு"
 SHUTDOWN     : "shutdown பண்ணு" / "restart பண்ணு"
 HELP         : "help சொல்லு" / "what can you do"
 EXIT         : "bye jarvis" / "exit பண்ணு"
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
"""

# ─── Standard Library ───────────────────────────────────────────────────────
import os, sys, io, time, datetime, subprocess, webbrowser
import tempfile, ctypes, json, re, threading

# ─── Audio Recording ────────────────────────────────────────────────────────
import sounddevice as sd
import soundfile   as sf
import numpy       as np

# ─── Speech Recognition ─────────────────────────────────────────────────────
import speech_recognition as sr

# ─── Tanglish TTS (Google TTS) ──────────────────────────────────────────────
from gtts      import gTTS
from playsound import playsound

# ─── Windows Control ────────────────────────────────────────────────────────
import pyautogui
import psutil
import pygetwindow as gw
from   PIL import ImageGrab

# ─── File Creation ──────────────────────────────────────────────────────────
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

# ─── WhatsApp ───────────────────────────────────────────────────────────────
try:
    import pywhatkit as pwk
    WHATSAPP_AVAILABLE = True
except ImportError:
    WHATSAPP_AVAILABLE = False

# ─── Volume (pycaw) ─────────────────────────────────────────────────────────
try:
    from pycaw.pycaw import AudioUtilities, IAudioEndpointVolume
    from comtypes   import CLSCTX_ALL
    VOLUME_AVAILABLE = True
except ImportError:
    VOLUME_AVAILABLE = False

# ─── Brightness ─────────────────────────────────────────────────────────────
try:
    import screen_brightness_control as sbc
    BRIGHTNESS_AVAILABLE = True
except ImportError:
    BRIGHTNESS_AVAILABLE = False

# ─── Win32 ──────────────────────────────────────────────────────────────────
try:
    import win32api, win32con, win32gui
    WIN32_AVAILABLE = True
except ImportError:
    WIN32_AVAILABLE = False

pyautogui.FAILSAFE = False

# ════════════════════════════════════════════════════════════════════════════
#  CONTACTS  (edit contacts.json to add your real numbers)
# ════════════════════════════════════════════════════════════════════════════
CONTACTS_FILE = os.path.join(os.path.dirname(__file__), "contacts.json")

def load_contacts() -> dict:
    try:
        with open(CONTACTS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}

CONTACTS = load_contacts()

# ════════════════════════════════════════════════════════════════════════════
#  TANGLISH TEXT-TO-SPEECH
#  Responses are written in Tamil+English mix — spoken via gTTS Tamil voice
# ════════════════════════════════════════════════════════════════════════════
def speak(text: str):
    """Speak Tanglish text using Google TTS (Tamil voice)."""
    print(f"\n  🤖 JARVIS: {text}")
    tmp = None
    try:
        tts = gTTS(text=text, lang='ta', slow=False)
        with tempfile.NamedTemporaryFile(delete=False, suffix='.mp3') as f:
            tmp = f.name
        tts.save(tmp)
        playsound(tmp)
    except Exception as e:
        print(f"  [TTS Error: {e}]")
    finally:
        if tmp:
            try: os.remove(tmp)
            except: pass


# ════════════════════════════════════════════════════════════════════════════
#  BILINGUAL MICROPHONE → GOOGLE SPEECH
#  Tries Tamil first, then English — handles Tanglish naturally
# ════════════════════════════════════════════════════════════════════════════
SAMPLE_RATE = 16000
CHANNELS    = 1
RECORD_SECS = 6

def record_audio() -> bytes:
    print("\n  🎤 Listening... (Tamil or English-ல பேசுங்க)")
    data = sd.rec(int(RECORD_SECS * SAMPLE_RATE),
                  samplerate=SAMPLE_RATE, channels=CHANNELS, dtype='int16')
    sd.wait()
    buf = io.BytesIO()
    sf.write(buf, data, SAMPLE_RATE, format='WAV', subtype='PCM_16')
    buf.seek(0)
    return buf.read()

def listen() -> str:
    """
    Record and transcribe in Tanglish.
    Tries Tamil (ta-IN) first — if nothing heard, tries English (en-IN).
    """
    recognizer = sr.Recognizer()
    wav_bytes   = record_audio()
    audio       = sr.AudioData(wav_bytes, SAMPLE_RATE, 2)

    # Try Tamil first
    for lang_code in ['ta-IN', 'en-IN']:
        try:
            cmd = recognizer.recognize_google(audio, language=lang_code).lower()
            if cmd.strip():
                lang_label = "Tamil" if lang_code == 'ta-IN' else "English"
                print(f"  👤 [{lang_label}] நீங்கள் சொன்னது: \"{cmd}\"")
                return cmd
        except sr.UnknownValueError:
            continue
        except sr.RequestError:
            speak("Network error da. Internet check பண்ணுங்க!")
            return ""

    speak("சரியா கேக்கல. மீண்டும் try பண்ணுங்க!")
    return ""


# ════════════════════════════════════════════════════════════════════════════
#  HELPERS
# ════════════════════════════════════════════════════════════════════════════
def has(text: str, *kw) -> bool:
    return any(k in text for k in kw)

def desktop_path(filename: str) -> str:
    return os.path.join(os.path.expanduser("~"), "Desktop", filename)


# ════════════════════════════════════════════════════════════════════════════
#  WINDOWS CONTROL FUNCTIONS
# ════════════════════════════════════════════════════════════════════════════

# ── VOLUME ──────────────────────────────────────────────────────────────────
def _vol_iface():
    if not VOLUME_AVAILABLE: return None
    try:
        dev = AudioUtilities.GetSpeakers()
        iface = dev.Activate(IAudioEndpointVolume._iid_, CLSCTX_ALL, None)
        return ctypes.cast(iface, ctypes.POINTER(IAudioEndpointVolume))
    except: return None

def volume_up(step=10):
    v = _vol_iface()
    if v:
        n = min(100, int(v.GetMasterVolumeLevelScalar()*100) + step)
        v.SetMasterVolumeLevelScalar(n/100, None)
        speak(f"Volume {n} percent-ஆ set பண்ணிட்டேன்!")
    else:
        pyautogui.press('volumeup', presses=5)
        speak("Volume up பண்ணிட்டேன்!")

def volume_down(step=10):
    v = _vol_iface()
    if v:
        n = max(0, int(v.GetMasterVolumeLevelScalar()*100) - step)
        v.SetMasterVolumeLevelScalar(n/100, None)
        speak(f"Volume {n} percent-க்கு குறைச்சிட்டேன்!")
    else:
        pyautogui.press('volumedown', presses=5)
        speak("Volume down பண்ணிட்டேன்!")

def volume_mute_toggle():
    v = _vol_iface()
    if v:
        muted = v.GetMute()
        v.SetMute(not muted, None)
        speak("Volume mute பண்ணிட்டேன்!" if not muted else "Volume unmute பண்ணிட்டேன்!")
    else:
        pyautogui.press('volumemute')
        speak("Volume mute toggle பண்ணிட்டேன்!")

def volume_set(percent: int):
    v = _vol_iface()
    if v:
        v.SetMasterVolumeLevelScalar(max(0,min(100,percent))/100, None)
        speak(f"Volume exactly {percent} percent-ஆ set பண்ணிட்டேன்!")
    else:
        speak("Volume control available இல்ல. Manual-ஆ try பண்ணுங்க!")

def volume_status():
    v = _vol_iface()
    if v:
        lvl   = int(v.GetMasterVolumeLevelScalar()*100)
        muted = v.GetMute()
        speak(f"இப்போ volume {lvl} percent. {'Mute-ல இருக்கு!' if muted else 'Sound on-ல இருக்கு!'}")
    else:
        speak("Volume info கிடைக்கல!")

# ── BRIGHTNESS ──────────────────────────────────────────────────────────────
def brightness_up(step=10):
    if BRIGHTNESS_AVAILABLE:
        try:
            cur = sbc.get_brightness(display=0)[0]
            new = min(100, cur + step)
            sbc.set_brightness(new, display=0)
            speak(f"Brightness {new} percent-ஆ increase பண்ணிட்டேன்!")
            return
        except: pass
    pyautogui.hotkey('fn', 'f6')
    speak("Brightness up பண்ணிட்டேன்!")

def brightness_down(step=10):
    if BRIGHTNESS_AVAILABLE:
        try:
            cur = sbc.get_brightness(display=0)[0]
            new = max(10, cur - step)
            sbc.set_brightness(new, display=0)
            speak(f"Brightness {new} percent-க்கு குறைச்சிட்டேன்!")
            return
        except: pass
    pyautogui.hotkey('fn', 'f5')
    speak("Brightness down பண்ணிட்டேன்!")

# ── SCREENSHOT ──────────────────────────────────────────────────────────────
def take_screenshot():
    ts   = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = desktop_path(f"Jarvis_Screenshot_{ts}.png")
    ImageGrab.grab().save(path)
    speak(f"Screenshot எடுத்திட்டேன்! Desktop-ல save ஆச்சு.")
    print(f"  📸 {path}")

# ── BATTERY ─────────────────────────────────────────────────────────────────
def battery_status():
    b = psutil.sensors_battery()
    if b:
        plugged = "charge ஆகுது" if b.power_plugged else "charge ஆகல"
        speak(f"Battery {int(b.percent)} percent இருக்கு. {plugged}.")
    else:
        speak("Battery info இல்ல. Desktop PC-ஆ இருக்கும்!")

# ── SYSTEM INFO ─────────────────────────────────────────────────────────────
def system_info():
    cpu = psutil.cpu_percent(interval=1)
    ram = psutil.virtual_memory()
    used  = round(ram.used/(1024**3), 1)
    total = round(ram.total/(1024**3), 1)
    speak(f"System status: CPU {cpu} percent use ஆகுது. RAM {used} GB use ஆகுது, total {total} GB இருக்கு.")

# ── WINDOW CONTROLS ─────────────────────────────────────────────────────────
def close_window():
    pyautogui.hotkey('alt','F4')
    speak("Window close பண்ணிட்டேன்!")

def minimize_window():
    pyautogui.hotkey('win','down')
    speak("Window minimize பண்ணிட்டேன்!")

def maximize_window():
    pyautogui.hotkey('win','up')
    speak("Window maximize பண்ணிட்டேன்!")

def switch_window():
    pyautogui.hotkey('alt','tab')
    speak("Window switch பண்ணிட்டேன்!")

def show_desktop():
    pyautogui.hotkey('win','d')
    speak("Desktop காட்டிட்டேன்!")

def show_all_windows():
    pyautogui.hotkey('win','tab')
    speak("All windows காட்டுறேன்!")

# ── LOCK / SLEEP ────────────────────────────────────────────────────────────
def lock_screen():
    if WIN32_AVAILABLE:
        win32api.LockWorkStation()
    else:
        os.system("rundll32.exe user32.dll,LockWorkStation")
    speak("Screen lock பண்ணிட்டேன்!")

def sleep_pc():
    speak("Computer sleep mode-க்கு போகுது. Bye!")
    time.sleep(1)
    os.system("rundll32.exe powrprof.dll,SetSuspendState 0,1,0")

# ── CLIPBOARD ───────────────────────────────────────────────────────────────
def copy_all():
    pyautogui.hotkey('ctrl','a')
    pyautogui.hotkey('ctrl','c')
    speak("Select all பண்ணி copy பண்ணிட்டேன்!")

def paste():
    pyautogui.hotkey('ctrl','v')
    speak("Paste பண்ணிட்டேன்!")

def undo():
    pyautogui.hotkey('ctrl','z')
    speak("Undo பண்ணிட்டேன்!")

# ── SCROLL ──────────────────────────────────────────────────────────────────
def scroll_down(): pyautogui.scroll(-5);  speak("Scroll down பண்ணிட்டேன்!")
def scroll_up():   pyautogui.scroll(5);   speak("Scroll up பண்ணிட்டேன்!")

# ── MEDIA ───────────────────────────────────────────────────────────────────
def media_play_pause(): pyautogui.press('playpause'); speak("Play pause பண்ணிட்டேன்!")
def media_next():       pyautogui.press('nexttrack'); speak("Next song போட்டேன்!")
def media_prev():       pyautogui.press('prevtrack'); speak("Previous song போட்டேன்!")

# ── TYPE TEXT ───────────────────────────────────────────────────────────────
def type_text(text: str):
    pyautogui.typewrite(text, interval=0.04)
    speak(f"Type பண்ணிட்டேன்!")

# ── WIFI ────────────────────────────────────────────────────────────────────
def wifi_on():
    os.system('netsh interface set interface "Wi-Fi" enable')
    speak("Wi-Fi enable பண்ணிட்டேன்!")

def wifi_off():
    os.system('netsh interface set interface "Wi-Fi" disable')
    speak("Wi-Fi disable பண்ணிட்டேன்!")

# ── RECYCLE BIN ─────────────────────────────────────────────────────────────
def empty_recycle_bin():
    os.system('PowerShell -Command "Clear-RecycleBin -Force" 2>nul')
    speak("Recycle bin empty பண்ணிட்டேன்!")

# ── WINDOWS UI ──────────────────────────────────────────────────────────────
def open_start_menu():    pyautogui.press('win');                speak("Start menu திறந்தேன்!")
def open_win_search():    pyautogui.hotkey('win','s');           speak("Windows search திறந்தேன்!")
def open_action_center(): pyautogui.hotkey('win','a');           speak("Action center திறந்தேன்!")
def open_settings():      pyautogui.hotkey('win','i');           speak("Settings திறந்தேன்!")
def open_run():           pyautogui.hotkey('win','r');           speak("Run dialog திறந்தேன்!")
def snap_left():          pyautogui.hotkey('win','left');        speak("Window-ஐ left-ல snap பண்ணிட்டேன்!")
def snap_right():         pyautogui.hotkey('win','right');       speak("Window-ஐ right-ல snap பண்ணிட்டேன்!")
def virtual_desktop_new():pyautogui.hotkey('win','ctrl','d');    speak("New virtual desktop create பண்ணிட்டேன்!")

def all_settings_on():
    """Turn on all Windows settings (Wi-Fi, Volume, Brightness, open Settings)."""
    speak("All settings turn on பண்றேன்!")
    os.system('netsh interface set interface "Wi-Fi" enable')
    v = _vol_iface()
    if v:
        v.SetMute(False, None)
        v.SetMasterVolumeLevelScalar(0.8, None)
    if BRIGHTNESS_AVAILABLE:
        try:
            sbc.set_brightness(80, display=0)
        except:
            pass
    pyautogui.hotkey('win','i')
    speak("All settings turned ON! Wi-Fi, Volume (80%), Brightness (80%) and Settings opened!")



# ════════════════════════════════════════════════════════════════════════════
#  FILE CREATION
# ════════════════════════════════════════════════════════════════════════════
def create_text_file(name: str, content: str = ""):
    """Create a .txt file on Desktop."""
    name = name.strip() or "Jarvis_Note"
    if not name.endswith(".txt"): name += ".txt"
    path = desktop_path(name)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content if content else f"# Created by Jarvis AI\n# Date: {datetime.datetime.now()}\n\n")
    speak(f"{name} text file create பண்ணிட்டேன்! Desktop-ல இருக்கு.")
    subprocess.Popen(["notepad.exe", path])

def create_word_file(name: str, content: str = ""):
    """Create a .docx Word file on Desktop."""
    if not DOCX_AVAILABLE:
        speak("Word file create பண்ண python-docx install பண்ணணும்!")
        return
    name = name.strip() or "Jarvis_Document"
    if not name.endswith(".docx"): name += ".docx"
    path = desktop_path(name)
    doc  = Document()
    doc.add_heading("Jarvis AI - Created Document", 0)
    doc.add_paragraph(f"Date: {datetime.datetime.now().strftime('%d %B %Y')}")
    doc.add_paragraph("")
    if content:
        doc.add_paragraph(content)
    doc.save(path)
    speak(f"{name} Word document create பண்ணிட்டேன்! Desktop-ல இருக்கு.")
    try: os.startfile(path)
    except: pass

def create_excel_file(name: str):
    """Create a .xlsx Excel file on Desktop."""
    if not XLSX_AVAILABLE:
        speak("Excel file create பண்ண openpyxl install பண்ணணும்!")
        return
    name = name.strip() or "Jarvis_Sheet"
    if not name.endswith(".xlsx"): name += ".xlsx"
    path = desktop_path(name)
    wb   = openpyxl.Workbook()
    ws   = wb.active
    ws.title = "Sheet1"
    ws['A1'] = "Created by Jarvis AI"
    ws['B1'] = str(datetime.datetime.now().strftime('%d %B %Y'))
    wb.save(path)
    speak(f"{name} Excel file create பண்ணிட்டேன்! Desktop-ல இருக்கு.")
    try: os.startfile(path)
    except: pass

def create_python_file(name: str):
    """Create a .py Python file on Desktop."""
    name = name.strip() or "jarvis_script"
    if not name.endswith(".py"): name += ".py"
    path = desktop_path(name)
    content = f'''# {name}
# Created by Jarvis AI - {datetime.datetime.now().strftime("%d %B %Y %H:%M")}

def main():
    print("Hello from {name}!")

if __name__ == "__main__":
    main()
'''
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    speak(f"{name} Python file create பண்ணிட்டேன்! Desktop-ல இருக்கு.")
    subprocess.Popen(["notepad.exe", path])

def create_folder(name: str):
    """Create a folder on Desktop."""
    name = name.strip() or "Jarvis_Folder"
    path = desktop_path(name)
    os.makedirs(path, exist_ok=True)
    speak(f"{name} folder create பண்ணிட்டேன்! Desktop-ல இருக்கு.")
    subprocess.Popen(["explorer.exe", path])

def handle_file_create(command: str):
    """
    Parse file creation commands like:
    "text file create பண்ணு MyNote"
    "word file create பண்ணு Report"
    "python file create பண்ணு script"
    "excel file create பண்ணு Data"
    "folder create பண்ணு Projects"
    """
    # Extract name after keywords
    name = ""
    for marker in ["பண்ணு ", "create ", "make ", "new "]:
        if marker in command:
            parts = command.split(marker)
            if len(parts) > 1:
                name = parts[-1].strip()
                break

    if has(command, "text file", "txt", "notepad file", "notes"):
        create_text_file(name)
    elif has(command, "word file", "word document", "doc file", "docx"):
        create_word_file(name)
    elif has(command, "excel file", "spreadsheet", "xlsx"):
        create_excel_file(name)
    elif has(command, "python file", "py file", "python script"):
        create_python_file(name)
    elif has(command, "folder", "directory"):
        create_folder(name)
    else:
        # Default: text file
        create_text_file(name or "JarvisNote")


# ════════════════════════════════════════════════════════════════════════════
#  WHATSAPP MESSAGING
# ════════════════════════════════════════════════════════════════════════════
def send_whatsapp(command: str):
    """
    Parse and send a WhatsApp message.
    Commands like:
      "whatsapp-ல அம்மாக்கு message அனுப்பு: Hello!"
      "send whatsapp to friend: Hi da"
      "whatsapp அம்மா: நான் வீட்டில் இருக்கேன்"
    """
    if not WHATSAPP_AVAILABLE:
        speak("pywhatkit install பண்ணல. WhatsApp send பண்ண முடியல!")
        return

    # Reload contacts (in case user edited file)
    contacts = load_contacts()

    # Extract contact name
    contact_name = None
    phone_number = None

    # Find who to send to
    for name, number in contacts.items():
        if name.lower() in command:
            contact_name = name
            phone_number = number
            break

    # Try to find inline phone number (+91XXXXXXXXXX)
    phone_match = re.search(r'\+\d{10,13}', command)
    if phone_match and not phone_number:
        phone_number = phone_match.group()
        contact_name = phone_number

    if not phone_number:
        speak(
            "யாருக்கு WhatsApp அனுப்பணும்? "
            "Contacts.json file-ல அவங்க பேரு add பண்ணுங்க. "
            "அல்லது phone number direct-ஆ சொல்லுங்க!"
        )
        return

    # Extract message (after colon or 'message' keyword)
    message = ""
    for sep in [":", "message ", "msg "]:
        if sep in command:
            parts = command.split(sep, 1)
            if len(parts) > 1:
                msg_part = parts[1].strip()
                # Remove contact name from message if accidentally included
                if contact_name:
                    msg_part = msg_part.replace(contact_name.lower(), "").strip()
                message = msg_part
                break

    if not message:
        speak("என்ன message அனுப்பணும்? Message content சொல்லுங்க!")
        return

    # Add Jarvis signature
    full_message = f"{message}\n\n- Jarvis AI மூலம் அனுப்பப்பட்டது"

    speak(f"{contact_name}-க்கு WhatsApp message அனுப்புறேன். WhatsApp Web-ல browser open ஆகும்!")

    try:
        now     = datetime.datetime.now()
        send_hr = now.hour
        send_mn = now.minute + 2   # 2 min from now
        if send_mn >= 60:
            send_hr += 1
            send_mn -= 60

        pwk.sendwhatmsg(
            phone_no      = phone_number,
            message       = full_message,
            time_hour     = send_hr,
            time_min      = send_mn,
            wait_time     = 20,
            tab_close     = True,
            close_time    = 5
        )
        speak(f"{contact_name}-க்கு message successfully அனுப்பிட்டேன்! Check பண்ணுங்க.")
    except Exception as e:
        speak(f"WhatsApp message அனுப்ப problem ஆச்சு. Error: {str(e)[:50]}")
        print(f"  [WhatsApp Error: {e}]")


# ════════════════════════════════════════════════════════════════════════════
#  APP LAUNCHER
# ════════════════════════════════════════════════════════════════════════════
APP_MAP = {
    "notepad": "notepad.exe",         "நோட்பேட்": "notepad.exe",
    "calculator": "calc.exe",         "கால்குலேட்டர்": "calc.exe",
    "paint": "mspaint.exe",           "பெயிண்ட்": "mspaint.exe",
    "cmd": "cmd.exe",                 "கட்டளை வரி": "cmd.exe",
    "command prompt": "cmd.exe",
    "task manager": "taskmgr.exe",    "பணி மேலாளர்": "taskmgr.exe",
    "file explorer": "explorer.exe",  "கோப்பு மேலாளர்": "explorer.exe",
    "file manager": "explorer.exe",
    "wordpad": "wordpad.exe",
    "snipping tool": "SnippingTool.exe",
    "control panel": "control.exe",
    "chrome":   r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "குரோம்":   r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "edge":     r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    "firefox":  r"C:\Program Files\Mozilla Firefox\firefox.exe",
    "word":     r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE",
    "excel":    r"C:\Program Files\Microsoft Office\root\Office16\EXCEL.EXE",
    "powerpoint": r"C:\Program Files\Microsoft Office\root\Office16\POWERPNT.EXE",
    "vlc":      r"C:\Program Files\VideoLAN\VLC\vlc.exe",
    "spotify":  r"%APPDATA%\Spotify\Spotify.exe",
    "vs code":  r"%LOCALAPPDATA%\Programs\Microsoft VS Code\Code.exe",
    "visual studio code": r"%LOCALAPPDATA%\Programs\Microsoft VS Code\Code.exe",
    "teams":    r"%LOCALAPPDATA%\Microsoft\Teams\current\Teams.exe",
    "zoom":     r"%APPDATA%\Zoom\bin\Zoom.exe",
    "telegram": r"%APPDATA%\Telegram Desktop\Telegram.exe",
}

def open_app(app_name: str) -> bool:
    path = APP_MAP.get(app_name.lower())
    if path is None:
        try: subprocess.Popen(app_name + ".exe", shell=True); return True
        except: return False
    if "%" in path: path = os.path.expandvars(path)
    if path.startswith("ms-"): os.system(f"start {path}"); return True
    if os.path.exists(path): subprocess.Popen([path]); return True
    try: subprocess.Popen(path, shell=True); return True
    except: return False


# ════════════════════════════════════════════════════════════════════════════
#  MASTER COMMAND PROCESSOR
# ════════════════════════════════════════════════════════════════════════════
def process_command(command: str) -> bool:
    """Route Tanglish command to correct action. Returns False to exit."""

    # ── EXIT ──────────────────────────────────────────────────────────────
    if has(command, "bye jarvis", "exit பண்ணு", "விடைபெறு", "goodbye",
           "shutdown jarvis", "close jarvis"):
        speak("Okay da! See you soon. Jarvis shutdown ஆகுது. Take care!")
        return False

    # ════ WHATSAPP ════════════════════════════════════════════════════════
    elif has(command, "whatsapp", "வாட்ஸ்அப்"):
        send_whatsapp(command)

    # ════ FILE CREATION ═══════════════════════════════════════════════════
    elif has(command, "file create", "create file", "make file", "new file",
             "file பண்ணு", "create பண்ணு", "folder create", "folder பண்ணு"):
        handle_file_create(command)

    # ════ VOLUME ══════════════════════════════════════════════════════════
    elif has(command, "volume up", "ஒலி அதிகரி", "sound up", "louder"):
        volume_up()
    elif has(command, "volume down", "ஒலி குறை", "sound down", "quieter"):
        volume_down()
    elif has(command, "mute", "முடக்கு") and has(command, "volume","sound","ஒலி"):
        volume_mute_toggle()
    elif has(command, "unmute", "volume on"):
        volume_mute_toggle()
    elif has(command, "volume status", "ஒலி நிலை"):
        volume_status()
    elif "volume" in command and any(c.isdigit() for c in command):
        # "volume 50 percent set பண்ணு"
        nums = re.findall(r'\d+', command)
        if nums: volume_set(int(nums[0]))

    # ════ BRIGHTNESS ══════════════════════════════════════════════════════
    elif has(command, "brightness up", "பிரகாசம் அதிகரி", "screen bright"):
        brightness_up()
    elif has(command, "brightness down", "பிரகாசம் குறை", "screen dim"):
        brightness_down()

    # ════ SCREENSHOT ══════════════════════════════════════════════════════
    elif has(command, "screenshot", "screen capture", "திரை படம்",
             "screen shot", "snap screen"):
        take_screenshot()

    # ════ BATTERY ═════════════════════════════════════════════════════════
    elif has(command, "battery", "பேட்டரி", "charge status"):
        battery_status()

    # ════ SYSTEM INFO ═════════════════════════════════════════════════════
    elif has(command, "system info", "cpu", "ram", "memory",
             "கணினி நிலை", "computer status"):
        system_info()

    # ════ WINDOW CONTROLS ═════════════════════════════════════════════════
    elif has(command, "close window", "window close", "சாளரம் மூடு", "alt f4"):
        close_window()
    elif has(command, "minimize", "window down", "சாளரம் குறைக்கவும்"):
        minimize_window()
    elif has(command, "maximize", "window up", "full screen", "சாளரம் பெரிதாக்கவும்"):
        maximize_window()
    elif has(command, "switch window", "alt tab", "சாளரம் மாற்று"):
        switch_window()
    elif has(command, "all windows", "task view", "அனைத்து சாளரம்"):
        show_all_windows()
    elif has(command, "snap left"):
        snap_left()
    elif has(command, "snap right"):
        snap_right()
    elif has(command, "new desktop", "virtual desktop"):
        virtual_desktop_new()

    # ════ DESKTOP ═════════════════════════════════════════════════════════
    elif has(command, "show desktop", "desktop காட்டு", "டெஸ்க்டாப்"):
        show_desktop()

    # ════ LOCK / SLEEP ════════════════════════════════════════════════════
    elif has(command, "lock screen", "screen lock", "திரை பூட்டு", "lock பண்ணு"):
        lock_screen()
    elif has(command, "sleep mode", "தூக்க நிலை", "sleep பண்ணு", "suspend"):
        sleep_pc()

    # ════ CLIPBOARD ═══════════════════════════════════════════════════════
    elif has(command, "copy", "நகலெடு"):
        copy_all()
    elif has(command, "paste", "ஒட்டு"):
        paste()
    elif has(command, "undo", "செயல் திரும்பு"):
        undo()

    # ════ SCROLL ══════════════════════════════════════════════════════════
    elif has(command, "scroll down", "கீழே உருட்டு"):
        scroll_down()
    elif has(command, "scroll up", "மேலே உருட்டு"):
        scroll_up()

    # ════ MEDIA ═══════════════════════════════════════════════════════════
    elif has(command, "play pause", "pause music", "pause song", "play music"):
        media_play_pause()
    elif has(command, "next song", "next track", "அடுத்த பாடல்"):
        media_next()
    elif has(command, "previous song", "prev song", "முந்தைய பாடல்"):
        media_prev()

    # ════ TYPE TEXT ═══════════════════════════════════════════════════════
    elif has(command, "type ", "தட்டச்சு"):
        text = command.replace("type","").replace("தட்டச்சு","").strip()
        if text: type_text(text)
        else: speak("என்ன type பண்ணணும்?")

    # ════ WINDOWS UI ══════════════════════════════════════════════════════
    elif has(command, "start menu", "தொடக்க மெனு"):
        open_start_menu()
    elif has(command, "windows search", "search bar", "win search"):
        open_win_search()
    elif has(command, "action center", "notification center"):
        open_action_center()
    elif has(command, "run dialog", "run command"):
        open_run()
    elif has(command, "settings open", "அமைப்புகள்") and not has(command, "open settings via"):
        open_settings()

    # ════ WIFI ════════════════════════════════════════════════════════════
    elif has(command, "wifi on", "wi-fi on", "enable wifi", "வைஃபை இயக்கு"):
        wifi_on()
    elif has(command, "wifi off", "wi-fi off", "disable wifi", "வைஃபை நிறுத்து"):
        wifi_off()

    # ════ RECYCLE BIN ═════════════════════════════════════════════════════
    elif has(command, "recycle bin", "குப்பை தொட்டி", "empty trash"):
        empty_recycle_bin()

    # ════ YOUTUBE ═════════════════════════════════════════════════════════
    elif has(command, "youtube", "யூட்யூப்"):
        speak("YouTube open பண்றேன்!")
        webbrowser.open("https://www.youtube.com")

    # ════ WEB SEARCH ══════════════════════════════════════════════════════
    elif has(command, "search", "google", "தேடு", "look up", "browse"):
        query = command
        for w in ["search","google","தேடு","look up","browse","in","on","ல"]:
            query = query.replace(w, "").strip()
        if query:
            speak(f"'{query}' Google-ல search பண்றேன்!")
            webbrowser.open(f"https://www.google.com/search?q={query.replace(' ','+')}")
        else:
            speak("என்ன search பண்ணணும்?")

    # ════ OPEN APP ════════════════════════════════════════════════════════
    elif has(command, "open", "திற", "launch", "start", "ஆரம்பி"):
        opened = False
        for key in APP_MAP:
            if key in command:
                speak(f"{key} open பண்றேன், one second!")
                if not open_app(key):
                    speak(f"Sorry da, {key} கண்டுபிடிக்கலை!")
                opened = True
                break
        if not opened:
            for trigger in ["open ", "launch ", "start "]:
                if trigger in command:
                    app = command.split(trigger,1)[1].strip()
                    if app:
                        speak(f"{app} open பண்றேன்!")
                        if not open_app(app):
                            speak(f"Sorry, {app} system-ல இல்ல மாதிரி!")
                        opened = True
                        break
        if not opened:
            speak("எந்த app open பண்ணணும்?")

    # ════ ALL SETTINGS ON ═════════════════════════════════════════════════
    elif has(command, "all settings on", "turn on all settings", "all setting on", "all settings"):
        all_settings_on()

    # ════ SETTINGS (via app open) ═════════════════════════════════════════
    elif has(command, "settings"):
        open_settings()

    # ════ TIME ════════════════════════════════════════════════════════════
    elif has(command, "time", "நேரம்", "மணி", "what time"):
        n   = datetime.datetime.now()
        per = "morning" if n.hour<12 else ("afternoon" if n.hour<17 else "evening")
        h12 = n.hour%12 or 12
        speak(f"இப்போ {per} {h12}:{n.minute:02d} ஆகுது!")

    # ════ DATE ════════════════════════════════════════════════════════════
    elif has(command, "date", "தேதி", "today", "இன்று"):
        t = datetime.datetime.now()
        months = {1:"January",2:"February",3:"March",4:"April",5:"May",
                  6:"June",7:"July",8:"August",9:"September",
                  10:"October",11:"November",12:"December"}
        speak(f"Today is {t.day} {months[t.month]} {t.year}!")

    # ════ SHUTDOWN PC ═════════════════════════════════════════════════════
    elif has(command, "shutdown", "கணினி அணை", "turn off computer"):
        speak("PC shutdown ஆகும். Cancel பண்ணணும்னா 'cancel' சொல்லுங்க!")
        time.sleep(2)
        c = listen()
        if has(c, "cancel", "ரத்து", "no", "வேண்டாம்"):
            speak("Okay, shutdown cancel பண்ணிட்டேன்!")
        else:
            speak("Shutting down now! Bye bye!")
            os.system("shutdown /s /t 1")
            return False

    # ════ RESTART ═════════════════════════════════════════════════════════
    elif has(command, "restart", "reboot", "மறுதொடக்கம்"):
        speak("PC restart ஆகும். Cancel-ஆ? 'cancel' சொல்லுங்க!")
        time.sleep(2)
        c = listen()
        if has(c, "cancel", "ரத்து", "no"):
            speak("Restart cancel பண்ணிட்டேன்!")
        else:
            speak("Restarting now!")
            os.system("shutdown /r /t 1")
            return False

    # ════ WHO ARE YOU ═════════════════════════════════════════════════════
    elif has(command, "who are you", "your name", "நீ யார்", "introduce"):
        speak(
            "Hey! நான் Jarvis. உங்கள் personal AI assistant. "
            "Tamil-ல பேசினாலும் English-ல பேசினாலும் mix பண்ணி பேசினாலும் "
            "நான் புரிஞ்சுக்குவேன்! "
            "Windows control, file create, WhatsApp message அனுப்பறது "
            "எல்லாமே நான் பண்ணுவேன். Just say the word!"
        )

    # ════ HELP ════════════════════════════════════════════════════════════
    elif has(command, "help", "உதவி", "what can you do", "commands", "jarvis"):
        speak(
            "Okay, listen carefully! "
            "Apps open பண்ணலாம். Volume, brightness control பண்ணலாம். "
            "Screenshot எடுக்கலாம். Battery, system info check பண்ணலாம். "
            "Windows close, minimize, maximize, snap பண்ணலாம். "
            "Screen lock, sleep mode போகலாம். Desktop காட்டலாம். "
            "Files create பண்ணலாம் - text, word, excel, python. "
            "WhatsApp message அனுப்பலாம். "
            "Google search, YouTube open பண்ணலாம். "
            "WiFi on off பண்ணலாம். Recycle bin empty பண்ணலாம். "
            "Media play pause, next, previous பண்ணலாம். "
            "Shutdown, restart பண்ணலாம். "
            "என்ன வேணும்னாலும் கேளுங்க! நான் ready!"
        )

    # ════ UNKNOWN ═════════════════════════════════════════════════════════
    else:
        speak(
            f"Hmm, '{command}' சரியா புரியல. "
            "Help அல்லது Jarvis-ன்னு சொல்லுங்க, all commands list பண்றேன்!"
        )

    return True  # Keep running


# ════════════════════════════════════════════════════════════════════════════
#  MAIN
# ════════════════════════════════════════════════════════════════════════════
def main():
    print("""
╔══════════════════════════════════════════════════════════════════════════╗
║    J A R V I S  AI  v4.0  -  Tanglish  Windows  Controller              ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Tamil + English Mix  |  All Windows Control  |  File Creator            ║
║  WhatsApp Messenger   |  Voice Activated       |  Always Listening        ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Say "help" or "jarvis" for all commands  |  Say "bye jarvis" to exit   ║
╚══════════════════════════════════════════════════════════════════════════╝
""")

    speak(
        "Hey! Jarvis here. Version 4 ready! "
        "Tamil-லயும் பேசலாம், English-லயும் பேசலாம், "
        "mix பண்ணியும் பேசலாம் — நான் எல்லாமே புரிஞ்சுக்குவேன்! "
        "Files create பண்ணலாம், WhatsApp அனுப்பலாம், "
        "Computer-ஐ full control பண்ணலாம். "
        "So, what do you need da?"
    )

    running = True
    while running:
        try:
            cmd = listen()
            if cmd:
                running = process_command(cmd)
        except KeyboardInterrupt:
            speak("Okay bye! Take care!")
            break

    print("\n  [Jarvis v4.0 stopped]")


if __name__ == "__main__":
    main()
